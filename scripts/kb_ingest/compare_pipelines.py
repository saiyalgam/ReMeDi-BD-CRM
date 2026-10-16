#!/usr/bin/env python3
"""
Compare two BusinessPipeline workbooks (read-only) and write a change report.

  python compare_pipelines.py --old "BusinessPipeline 10 July 2026.xlsx" --new "BusinessPipeline 29 Sep 2026.xlsx" --out <folder>

Writes pipeline_comparison.csv and pipeline_comparison.md. The workbooks are never modified.
Only the "Leads Tracking" sheet is compared. Contact sheets (names, phone numbers, e-mails) are counted but their
contents are NOT copied into the report, because they contain personal data.
Matching is by normalised project name; uncertain matches are reported as "possible match - confirm".
"""
import argparse, csv, difflib, os, re
import openpyxl

HEAD = ['S. No.', 'Name of the Project', 'Category', 'Size (Cr.)', 'Present State', 'Lead Owner', 'Geography', 'Details', 'Comments', 'Partner(s)']


def norm(s):
    s = (s or '').lower()
    s = re.sub(r'\(.*?\)|\b(inr|rs)\b[\d\s,.l]*|\b\d+(st|nd|rd|th)?\s*(aug|sep|sept|oct|nov|dec|jan)\w*', ' ', s)
    return re.sub(r'[^a-z0-9]+', ' ', s).strip()


def read(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb['Leads Tracking']
    rows = list(ws.iter_rows(values_only=True))
    out, section, hdr, last_hdr = [], 'Open leads', None, None
    for i, r in enumerate(rows, 1):
        vals = ['' if v is None else str(v).strip() for v in r]
        first = next((v for v in vals if v), '')
        if first in ('Won Leads', 'Lost Leads'):
            section = first; hdr = last_hdr; continue
        if 'Name of the Project' in vals:
            hdr = last_hdr = [v for v in vals]; continue
        if 'Classification by Present State' in vals:
            hdr = None; continue
        if hdr is None or not first or first.upper() == 'TOTAL':
            continue
        rec = {h: vals[j] if j < len(vals) else '' for j, h in enumerate(hdr) if h}
        name = rec.get('Name of the Project', '')
        if not name:
            # rows that only carry a serial number
            continue
        rec.update(_row=i, _section=section)
        out.append(rec)
    summary = []
    for i, r in enumerate(rows, 1):
        vals = ['' if v is None else str(v).strip() for v in r]
        if vals and 'Classification by Present State' in vals:
            summary.append(i)
    other = {ws2.title: ws2.max_row for ws2 in wb.worksheets if ws2.title != 'Leads Tracking'}
    return out, other


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--old', required=True); ap.add_argument('--new', required=True); ap.add_argument('--out', required=True)
    a = ap.parse_args()
    old, old_other = read(a.old)
    new, new_other = read(a.new)
    on = {norm(r['Name of the Project']): r for r in old}
    nn = {norm(r['Name of the Project']): r for r in new}
    rows, used_new = [], set()
    for k, r in on.items():
        if k in nn:
            m, how = nn[k], 'same name'
        else:
            cand = difflib.get_close_matches(k, list(nn), n=1, cutoff=0.72)
            m, how = (nn[cand[0]], 'possible match - confirm') if cand else (None, '')
        if m is None:
            rows.append(dict(change='only in older workbook', project=r['Name of the Project'], old_state=r.get('Present State', ''), new_state='',
                             old_owner=r.get('Lead Owner', ''), new_owner='', old_size=r.get('Size (Cr.)', ''), new_size='',
                             old_row=r['_row'], new_row='', note='Not carried into the newer workbook. Status unknown: confirm whether it was closed, dropped or forgotten.'))
            continue
        used_new.add(norm(m['Name of the Project']))
        diffs = [f for f in ('Present State', 'Lead Owner', 'Size (Cr.)', 'Category', 'Geography', 'Partner(s)') if (r.get(f) or '') != (m.get(f) or '')]
        rows.append(dict(change=('changed: ' + ', '.join(diffs)) if diffs else 'unchanged', match=how, project=r['Name of the Project'],
                         old_state=r.get('Present State', ''), new_state=m.get('Present State', '') + (f" ({m['_section']})" if m['_section'] != 'Open leads' else ''),
                         old_owner=r.get('Lead Owner', ''), new_owner=m.get('Lead Owner', ''), old_size=r.get('Size (Cr.)', ''), new_size=m.get('Size (Cr.)', ''),
                         old_row=r['_row'], new_row=m['_row'], note=''))
    for k, m in nn.items():
        if k not in used_new:
            rows.append(dict(change='only in newer workbook', project=m['Name of the Project'], old_state='', new_state=m.get('Present State', '') + (f" ({m['_section']})" if m['_section'] != 'Open leads' else ''),
                             old_owner='', new_owner=m.get('Lead Owner', ''), old_size='', new_size=m.get('Size (Cr.)', ''), old_row='', new_row=m['_row'], note=''))
    os.makedirs(a.out, exist_ok=True)
    cols = ['change', 'match', 'project', 'old_state', 'new_state', 'old_owner', 'new_owner', 'old_size', 'new_size', 'old_row', 'new_row', 'note']
    with open(os.path.join(a.out, 'pipeline_comparison.csv'), 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction='ignore'); w.writeheader(); [w.writerow(r) for r in rows]
    from collections import Counter
    c = Counter(r['change'].split(':')[0] for r in rows)
    md = [f'# BusinessPipeline comparison', '', f'- Older: `{os.path.basename(a.old)}` — {len(old)} named leads in "Leads Tracking"',
          f'- Newer: `{os.path.basename(a.new)}` — {len(new)} named leads in "Leads Tracking" (open, won and lost sections)', '']
    md += [f'- {k}: **{v}**' for k, v in c.most_common()]
    md += ['', '## Observations for a person to confirm', '',
           '- The two workbooks track largely **different sets** of leads (the older one is mostly foundation/CSR/award applications; the newer one is mostly government and overseas business). Most older leads do not appear in the newer file — that is not evidence that they were closed.',
           '- The status vocabulary changed between versions (for example the older summary uses "Tanked"; the newer summary uses "Banked"). Map these to CRM stages only after the team confirms what each means.',
           '- The "Size (Cr.)" column is mostly blank in the older file and filled in the newer one; amounts in the older file sometimes sit in other columns (e.g. "10,00,000", "INR 63L"). They are not converted here.',
           '- The newer file has a separate "Pre-Leads Tracking" sheet; the older file has contact sheets ("Personal Leads", "AVPN meetings"). Contact sheets hold names, phone numbers and e-mail addresses: they are personal data and were **not** copied into this report or the CRM knowledge base.',
           '', '## Sheets other than "Leads Tracking"', '', f'- Older: {old_other}', f'- Newer: {new_other}', '',
           '## Lead by lead', '', '| Change | Project | Older state | Newer state | Owner (old → new) | Size Cr (old → new) |', '|---|---|---|---|---|---|']
    md += [f"| {r['change']}{' (' + r['match'] + ')' if r.get('match') and r['match'] != 'same name' else ''} | {r['project']} | {r['old_state']} | {r['new_state']} | {r['old_owner']} → {r['new_owner']} | {r['old_size']} → {r['new_size']} |" for r in rows]
    open(os.path.join(a.out, 'pipeline_comparison.md'), 'w', encoding='utf-8').write('\n'.join(md) + '\n')
    print('\n'.join(md[:12]))


if __name__ == '__main__':
    main()
